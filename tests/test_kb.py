from __future__ import annotations

from datetime import datetime

import app.models  # noqa: F401
from app.config import Settings
from app.database import Base, SessionLocal, configure_database
from app.models import Announcement, Attachment, AttachmentVersion, SearchIndex, Site, SiteSection
from app.services.kb import (
    ask_knowledge,
    format_openclaw_answer,
    parse_attachment,
    parse_attachments,
    rebuild_search_index,
    search_knowledge,
)


def setup_db(tmp_path):
    engine = configure_database(f"sqlite:///{tmp_path / 'kb.db'}")
    Base.metadata.create_all(engine)


def seed_attachment(local_path: str, name: str = "测试附件.docx") -> int:
    with SessionLocal() as db:
        site = Site(
            name="住房城乡建设部",
            slug="mohurd",
            homepage_url="https://www.mohurd.gov.cn/",
            enabled=True,
        )
        db.add(site)
        db.flush()
        section = SiteSection(
            site_id=site.id,
            name="资质公告",
            url="https://www.mohurd.gov.cn/list.html",
            enabled=True,
        )
        db.add(section)
        db.flush()
        announcement = Announcement(
            site_id=site.id,
            section_id=section.id,
            identity_key="notice-1",
            identity_strategy="url",
            title="建筑业企业资质延续公告",
            item_type="qualification_notice",
            source_url="https://www.mohurd.gov.cn/a.html",
            content="本公告涉及建筑业企业资质延续和施工总承包。",
            published_at=datetime(2026, 6, 22, 9, 0),
        )
        db.add(announcement)
        db.flush()
        attachment = Attachment(
            announcement_id=announcement.id,
            site_id=site.id,
            attachment_key="att-1",
            name=name,
            safe_name=name,
            file_ext=f".{name.rsplit('.', 1)[-1]}",
            source_url="https://www.mohurd.gov.cn/a.docx",
            local_path=local_path,
            download_status="success",
        )
        db.add(attachment)
        db.flush()
        db.add(
            AttachmentVersion(
                attachment_id=attachment.id,
                announcement_id=announcement.id,
                site_id=site.id,
                version_no=1,
                name=attachment.name,
                safe_name=attachment.safe_name,
                source_url=attachment.source_url,
                local_path=local_path,
                file_hash="hash-1",
                change_type="new_attachment",
            )
        )
        db.commit()
        return attachment.id


def create_docx(path, text: str) -> None:
    from docx import Document

    document = Document()
    document.add_paragraph(text)
    document.save(path)


def create_xlsx(path, text: str) -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "公告名单"
    sheet["A1"] = text
    workbook.save(path)


def test_parse_docx_attachment_and_searches_text(tmp_path):
    setup_db(tmp_path)
    storage = tmp_path / "storage"
    file_path = storage / "attachments" / "notice.docx"
    file_path.parent.mkdir(parents=True)
    create_docx(file_path, "建筑业企业资质延续名单")
    attachment_id = seed_attachment("attachments/notice.docx")

    with SessionLocal() as db:
        document_text = parse_attachment(
            db,
            attachment_id,
            settings=Settings(APP_STORAGE_ROOT=storage),
        )
        assert document_text.status == "success"
        assert "资质延续名单" in document_text.text

        results = search_knowledge(db, "资质延续", limit=5)

    assert results
    assert results[0].entity_type == "attachment"


def test_parse_xlsx_attachment(tmp_path):
    setup_db(tmp_path)
    storage = tmp_path / "storage"
    file_path = storage / "attachments" / "notice.xlsx"
    file_path.parent.mkdir(parents=True)
    create_xlsx(file_path, "施工总承包二级")
    attachment_id = seed_attachment("attachments/notice.xlsx", name="名单.xlsx")

    with SessionLocal() as db:
        document_text = parse_attachment(
            db,
            attachment_id,
            settings=Settings(APP_STORAGE_ROOT=storage),
        )

    assert document_text.status == "success"
    assert "施工总承包二级" in document_text.text


def test_parse_attachments_records_missing_file_failure(tmp_path):
    setup_db(tmp_path)
    storage = tmp_path / "storage"
    attachment_id = seed_attachment("attachments/missing.docx")

    with SessionLocal() as db:
        summary = parse_attachments(db, limit=10, settings=Settings(APP_STORAGE_ROOT=storage))
        document_text = parse_attachment(
            db,
            attachment_id,
            settings=Settings(APP_STORAGE_ROOT=storage),
        )

    assert summary.total == 1
    assert summary.failed == 1
    assert document_text.status == "failed"
    assert "本地文件不存在" in document_text.error_message


def test_rebuild_search_index_and_ask(tmp_path):
    setup_db(tmp_path)
    storage = tmp_path / "storage"
    file_path = storage / "attachments" / "notice.docx"
    file_path.parent.mkdir(parents=True)
    create_docx(file_path, "资质核准公告附件内容")
    attachment_id = seed_attachment("attachments/notice.docx")

    with SessionLocal() as db:
        parse_attachment(db, attachment_id, settings=Settings(APP_STORAGE_ROOT=storage))
        count = rebuild_search_index(db)
        results = search_knowledge(db, "建筑业企业资质", limit=5)
        answer = ask_knowledge(db, "资质核准", limit=5, base_url="http://localhost:8000")
        openclaw_text = format_openclaw_answer(
            db,
            "资质核准",
            limit=5,
            base_url="http://localhost:8000",
        )

    assert count >= 2
    assert any(item.entity_type == "announcement" for item in results)
    assert answer["answer_type"] == "search_summary"
    assert answer["items"]
    assert answer["items"][0]["backend_url"].startswith("http://localhost:8000/")
    assert "后台：http://localhost:8000/" in openclaw_text


def test_search_knowledge_filters_by_entity_type_and_site(tmp_path):
    setup_db(tmp_path)
    with SessionLocal() as db:
        db.add_all(
            [
                SearchIndex(
                    entity_type="announcement",
                    entity_id=1,
                    title="建筑业企业资质延续公告",
                    body="建筑业企业资质延续",
                    site_name="住房城乡建设部",
                    section_name="资质公告",
                    published_at=datetime(2026, 6, 22, 9, 0),
                ),
                SearchIndex(
                    entity_type="attachment",
                    entity_id=2,
                    title="陕西资质延续名单.xlsx",
                    body="建筑业企业资质延续",
                    site_name="陕西省住房和城乡建设厅",
                    section_name="公告公示",
                    published_at=datetime(2026, 6, 24, 9, 0),
                ),
            ]
        )
        db.commit()

        announcement_results = search_knowledge(
            db,
            "资质延续",
            entity_type="announcement",
        )
        site_results = search_knowledge(
            db,
            "资质延续",
            site_name="陕西省住房和城乡建设厅",
        )
        date_results = search_knowledge(
            db,
            "资质延续",
            published_from=datetime(2026, 6, 23),
            published_to=datetime(2026, 6, 25),
        )

    assert [item.entity_type for item in announcement_results] == ["announcement"]
    assert [item.site_name for item in site_results] == ["陕西省住房和城乡建设厅"]
    assert [item.title for item in date_results] == ["陕西资质延续名单.xlsx"]
