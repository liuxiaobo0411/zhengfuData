from __future__ import annotations

import warnings
from dataclasses import dataclass
from datetime import datetime

from openpyxl import Workbook
from sqlalchemy import inspect, select
from sqlalchemy.orm import Session
from starlette.exceptions import StarletteDeprecationWarning

with warnings.catch_warnings():
    warnings.filterwarnings("ignore", category=StarletteDeprecationWarning)
    from fastapi.testclient import TestClient

from app.config import Settings
from app.main import create_app
from app.models import Attachment
from app.services.kb import (
    ask_knowledge,
    extract_docx_text,
    extract_xlsx_text,
    kb_stats,
    parse_attachment,
    rebuild_search_index,
    search_knowledge,
)
from app.services.storage import prepare_storage


@dataclass(frozen=True)
class V2AcceptanceCheck:
    name: str
    status: str
    detail: str


@dataclass(frozen=True)
class V2AcceptanceReport:
    checks: list[V2AcceptanceCheck]
    path: str | None = None

    @property
    def failed_count(self) -> int:
        return sum(1 for check in self.checks if check.status == "fail")

    @property
    def passed_count(self) -> int:
        return sum(1 for check in self.checks if check.status == "ok")


def run_v2_acceptance_check(db: Session, settings: Settings) -> V2AcceptanceReport:
    checks: list[V2AcceptanceCheck] = []
    checks.append(check_tables(db))
    checks.append(check_pdf_parse(db, settings))
    checks.append(check_docx_parser(settings))
    checks.append(check_xlsx_parser(settings))
    checks.append(check_search_index(db))
    checks.append(check_search_and_ask(db))
    checks.extend(check_api_and_web(settings))
    report = V2AcceptanceReport(checks=checks)
    path = write_v2_acceptance_report(report, settings)
    return V2AcceptanceReport(checks=checks, path=str(path))


def check_tables(db: Session) -> V2AcceptanceCheck:
    tables = set(inspect(db.bind).get_table_names())
    missing = {"document_texts", "search_index"} - tables
    if missing:
        return V2AcceptanceCheck("tables", "fail", f"缺少表：{', '.join(sorted(missing))}")
    return V2AcceptanceCheck("tables", "ok", "document_texts 和 search_index 已存在")


def check_pdf_parse(db: Session, settings: Settings) -> V2AcceptanceCheck:
    attachment = db.scalar(
        select(Attachment)
        .where(Attachment.local_path.is_not(None))
        .where(Attachment.download_status == "success")
        .where(Attachment.file_ext == ".pdf")
        .order_by(Attachment.id)
    )
    if attachment is None:
        return V2AcceptanceCheck("pdf_parse", "fail", "没有可用于验收的本地 PDF 附件")
    document_text = parse_attachment(db, attachment.id, settings=settings)
    if document_text.status == "success" and document_text.text_length > 0:
        return V2AcceptanceCheck(
            "pdf_parse",
            "ok",
            f"attachment={attachment.id} text_length={document_text.text_length}",
        )
    return V2AcceptanceCheck(
        "pdf_parse",
        "fail",
        (
            f"attachment={attachment.id} status={document_text.status} "
            f"reason={document_text.error_message}"
        ),
    )


def check_docx_parser(settings: Settings) -> V2AcceptanceCheck:
    from docx import Document

    sample_dir = prepare_storage(settings).exports / "v2_acceptance_samples"
    sample_dir.mkdir(parents=True, exist_ok=True)
    path = sample_dir / "sample.docx"
    document = Document()
    document.add_paragraph("V2 DOCX 解析样例：建筑业企业资质延续")
    document.save(path)
    text = extract_docx_text(path)
    if "建筑业企业资质延续" in text:
        return V2AcceptanceCheck("docx_parse", "ok", f"text_length={len(text)}")
    return V2AcceptanceCheck("docx_parse", "fail", "DOCX 样例未提取到预期文本")


def check_xlsx_parser(settings: Settings) -> V2AcceptanceCheck:
    sample_dir = prepare_storage(settings).exports / "v2_acceptance_samples"
    sample_dir.mkdir(parents=True, exist_ok=True)
    path = sample_dir / "sample.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "资质名单"
    sheet["A1"] = "V2 XLSX 解析样例：施工总承包二级"
    workbook.save(path)
    text = extract_xlsx_text(path)
    if "施工总承包二级" in text:
        return V2AcceptanceCheck("xlsx_parse", "ok", f"text_length={len(text)}")
    return V2AcceptanceCheck("xlsx_parse", "fail", "XLSX 样例未提取到预期文本")


def check_search_index(db: Session) -> V2AcceptanceCheck:
    count = rebuild_search_index(db)
    stats = kb_stats(db)
    if stats["search_index"] > 0:
        return V2AcceptanceCheck(
            "search_index",
            "ok",
            f"rebuilt={count} total={stats['search_index']}",
        )
    return V2AcceptanceCheck("search_index", "fail", "索引为空")


def check_search_and_ask(db: Session) -> V2AcceptanceCheck:
    results = search_knowledge(db, "资质", limit=5)
    answer = ask_knowledge(db, "资质", limit=5)
    if results and answer["items"]:
        return V2AcceptanceCheck("search_ask", "ok", f"results={len(results)}")
    return V2AcceptanceCheck("search_ask", "fail", "检索或问答摘要没有返回结果")


def check_api_and_web(settings: Settings) -> list[V2AcceptanceCheck]:
    client = TestClient(create_app())
    checks: list[V2AcceptanceCheck] = []
    search_response = client.post("/api/kb/search", json={"query": "资质", "limit": 5})
    checks.append(status_check("api_kb_search", search_response.status_code, 200))
    filtered_search_response = client.post(
        "/api/kb/search",
        json={"query": "资质", "limit": 5, "entity_type": "announcement"},
    )
    checks.append(status_check("api_kb_search_filters", filtered_search_response.status_code, 200))
    invalid_filter_response = client.post(
        "/api/kb/search",
        json={"query": "资质", "entity_type": "invalid"},
    )
    checks.append(
        status_check("api_kb_search_invalid_filter", invalid_filter_response.status_code, 422)
    )
    ask_response = client.post("/api/kb/ask", json={"question": "资质", "limit": 5})
    checks.append(status_check("api_kb_ask", ask_response.status_code, 200))
    openclaw_response = client.post("/api/openclaw/kb/ask", json={"text": "资质", "limit": 5})
    checks.append(status_check("api_openclaw_kb_ask", openclaw_response.status_code, 200))
    openclaw_text = ""
    if openclaw_response.status_code == 200:
        openclaw_text = openclaw_response.json().get("text", "")
    if settings.app_public_base_url in openclaw_text:
        checks.append(V2AcceptanceCheck("api_openclaw_backend_url", "ok", "返回完整后台链接"))
    else:
        checks.append(V2AcceptanceCheck("api_openclaw_backend_url", "fail", "未返回完整后台链接"))
    client.post(
        "/login",
        data={"username": settings.admin_username, "password": settings.admin_password},
        follow_redirects=False,
    )
    web_response = client.get("/kb/search?q=资质")
    checks.append(status_check("web_kb_search", web_response.status_code, 200))
    return checks


def status_check(name: str, actual: int, expected: int) -> V2AcceptanceCheck:
    if actual == expected:
        return V2AcceptanceCheck(name, "ok", f"status={actual}")
    return V2AcceptanceCheck(name, "fail", f"status={actual} expected={expected}")


def format_v2_acceptance_report(report: V2AcceptanceReport) -> str:
    lines = ["V2 验收结果"]
    for check in report.checks:
        label = "OK" if check.status == "ok" else "FAIL"
        lines.append(f"[{label}] {check.name}: {check.detail}")
    lines.append(f"summary ok={report.passed_count} fail={report.failed_count}")
    if report.path:
        lines.append(f"v2_acceptance_report={report.path}")
    return "\n".join(lines)


def write_v2_acceptance_report(report: V2AcceptanceReport, settings: Settings):
    exports = prepare_storage(settings).exports
    path = exports / f"v2_acceptance_report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.md"
    lines = [
        "# V2 验收报告",
        "",
        f"生成时间：{datetime.now().isoformat(timespec='seconds')}",
        "",
        "| 检查项 | 状态 | 详情 |",
        "| --- | --- | --- |",
    ]
    for check in report.checks:
        lines.append(f"| {check.name} | {check.status} | {check.detail} |")
    lines.extend(
        [
            "",
            f"汇总：ok={report.passed_count} fail={report.failed_count}",
            "",
        ]
    )
    path.write_text("\n".join(lines), encoding="utf-8")
    return path
