from __future__ import annotations

import hashlib
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader
from sqlalchemy import delete, func, or_, select
from sqlalchemy.orm import Session

from app.config import Settings, get_settings
from app.models import (
    Announcement,
    Attachment,
    AttachmentVersion,
    DocumentText,
    SearchIndex,
    Site,
    SiteSection,
)

PARSER_VERSION = "v1"
SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".xlsx", ".xlsm"}
SEARCH_ENTITY_TYPES = {"announcement", "attachment"}


@dataclass(frozen=True)
class ParseSummary:
    total: int
    success: int
    failed: int
    unsupported: int


@dataclass(frozen=True)
class KnowledgeSearchResult:
    entity_type: str
    entity_id: int
    title: str
    snippet: str
    match_fields: list[str]
    source_url: str | None
    backend_path: str | None
    site_name: str | None
    section_name: str | None
    published_at: datetime | None
    score: int


def parse_attachment(
    db: Session,
    attachment_id: int,
    settings: Settings | None = None,
) -> DocumentText:
    settings = settings or get_settings()
    attachment = db.get(Attachment, attachment_id)
    if attachment is None:
        raise ValueError(f"attachment not found: {attachment_id}")

    version = latest_attachment_version(db, attachment.id)
    document_text = get_or_create_document_text(db, attachment, version)
    document_text.status = "running"
    document_text.error_message = None
    db.flush()

    file_path = resolve_attachment_file(settings, attachment, version)
    suffix = file_path.suffix.lower() or (attachment.file_ext or "").lower()
    parser_name = parser_name_for_suffix(suffix)
    document_text.parser_name = parser_name
    document_text.parser_version = PARSER_VERSION

    if suffix not in SUPPORTED_EXTENSIONS:
        document_text.status = "unsupported"
        document_text.text = None
        document_text.text_hash = None
        document_text.text_length = 0
        document_text.error_message = f"不支持的附件格式：{suffix or 'unknown'}"
        document_text.parsed_at = datetime.now()
        db.commit()
        db.refresh(document_text)
        return document_text

    if not file_path.exists():
        document_text.status = "failed"
        document_text.text = None
        document_text.text_hash = None
        document_text.text_length = 0
        document_text.error_message = f"本地文件不存在：{file_path}"
        document_text.parsed_at = datetime.now()
        db.commit()
        db.refresh(document_text)
        return document_text

    try:
        text = extract_text(file_path, suffix)
        document_text.status = "success"
        document_text.text = text
        document_text.text_hash = sha256_text(text)
        document_text.text_length = len(text)
        document_text.error_message = None
    except Exception as exc:
        document_text.status = "failed"
        document_text.text = None
        document_text.text_hash = None
        document_text.text_length = 0
        document_text.error_message = f"{type(exc).__name__}: {exc}"
    document_text.parsed_at = datetime.now()
    db.commit()
    db.refresh(document_text)

    if document_text.status == "success":
        upsert_attachment_index(db, attachment, document_text)
    return document_text


def parse_attachments(
    db: Session,
    limit: int = 20,
    settings: Settings | None = None,
) -> ParseSummary:
    settings = settings or get_settings()
    query = (
        select(Attachment)
        .outerjoin(DocumentText, DocumentText.attachment_id == Attachment.id)
        .where(Attachment.local_path.is_not(None))
        .where(Attachment.download_status == "success")
        .where(or_(DocumentText.id.is_(None), DocumentText.status.in_(["failed", "unsupported"])))
        .order_by(Attachment.id)
    )
    if limit > 0:
        query = query.limit(limit)
    attachments = db.scalars(query).all()
    counts = {"success": 0, "failed": 0, "unsupported": 0}
    for attachment in attachments:
        document_text = parse_attachment(db, attachment.id, settings=settings)
        if document_text.status in counts:
            counts[document_text.status] += 1
        elif document_text.status != "success":
            counts["failed"] += 1
    return ParseSummary(
        total=len(attachments),
        success=counts["success"],
        failed=counts["failed"],
        unsupported=counts["unsupported"],
    )


def rebuild_search_index(db: Session) -> int:
    db.execute(delete(SearchIndex))
    db.flush()
    count = 0
    announcements = db.scalars(select(Announcement).order_by(Announcement.id)).all()
    for announcement in announcements:
        upsert_announcement_index(db, announcement, commit=False)
        count += 1
    document_texts = db.scalars(
        select(DocumentText).where(DocumentText.status == "success").order_by(DocumentText.id)
    ).all()
    for document_text in document_texts:
        attachment = db.get(Attachment, document_text.attachment_id)
        if attachment is None:
            continue
        upsert_attachment_index(db, attachment, document_text, commit=False)
        count += 1
    db.commit()
    return count


def upsert_announcement_index(
    db: Session,
    announcement: Announcement,
    commit: bool = True,
) -> SearchIndex:
    site = db.get(Site, announcement.site_id)
    section = db.get(SiteSection, announcement.section_id)
    body = "\n".join(
        part
        for part in [announcement.title, announcement.content_summary, announcement.content]
        if part
    )
    item = get_or_create_search_index(db, "announcement", announcement.id)
    item.title = announcement.title
    item.body = body
    item.source_url = announcement.source_url
    item.backend_path = f"/announcements/{announcement.id}"
    item.site_name = site.name if site else None
    item.section_name = section.name if section else None
    item.published_at = announcement.published_at
    item.indexed_at = datetime.now()
    if commit:
        db.commit()
        db.refresh(item)
    return item


def upsert_attachment_index(
    db: Session,
    attachment: Attachment,
    document_text: DocumentText,
    commit: bool = True,
) -> SearchIndex:
    announcement = db.get(Announcement, attachment.announcement_id)
    site = db.get(Site, attachment.site_id)
    section = db.get(SiteSection, announcement.section_id) if announcement else None
    body = "\n".join(part for part in [attachment.name, document_text.text] if part)
    item = get_or_create_search_index(db, "attachment", attachment.id)
    item.title = attachment.name
    item.body = body
    item.source_url = attachment.source_url
    item.backend_path = f"/attachments/{attachment.id}/download"
    item.site_name = site.name if site else None
    item.section_name = section.name if section else None
    item.published_at = announcement.published_at if announcement else None
    item.indexed_at = datetime.now()
    if commit:
        db.commit()
        db.refresh(item)
    return item


def search_knowledge(
    db: Session,
    query: str,
    limit: int = 10,
    entity_type: str | None = None,
    site_name: str | None = None,
    published_from: datetime | None = None,
    published_to: datetime | None = None,
) -> list[KnowledgeSearchResult]:
    terms = [term.strip() for term in query.split() if term.strip()]
    if not terms and query.strip():
        terms = [query.strip()]
    if not terms:
        return []

    statement = select(SearchIndex)
    if entity_type:
        statement = statement.where(SearchIndex.entity_type == entity_type)
    if site_name:
        statement = statement.where(SearchIndex.site_name == site_name)
    if published_from:
        statement = statement.where(SearchIndex.published_at >= published_from)
    if published_to:
        statement = statement.where(SearchIndex.published_at <= published_to)
    for term in terms:
        pattern = f"%{term}%"
        statement = statement.where(
            or_(
                SearchIndex.title.like(pattern),
                SearchIndex.body.like(pattern),
                SearchIndex.site_name.like(pattern),
                SearchIndex.section_name.like(pattern),
            )
        )
    rows = db.scalars(
        statement.order_by(SearchIndex.published_at.desc(), SearchIndex.id.desc()).limit(limit * 5)
    ).all()
    results = [result_from_index(row, terms) for row in rows]
    results.sort(key=lambda item: (item.score, item.published_at or datetime.min), reverse=True)
    return results[:limit]


def ask_knowledge(
    db: Session,
    question: str,
    limit: int = 5,
    base_url: str | None = None,
) -> dict:
    items = search_knowledge(db, question, limit=limit)
    answer = f"找到 {len(items)} 条相关信息，请以原文和附件为准。"
    return {
        "question": question,
        "answer_type": "search_summary",
        "answer": answer,
        "items": [search_result_to_dict(item, base_url=base_url) for item in items],
    }


def format_openclaw_answer(
    db: Session,
    text: str,
    limit: int = 5,
    base_url: str | None = None,
) -> str:
    result = ask_knowledge(db, text, limit=limit, base_url=base_url)
    items = result["items"]
    if not items:
        return f"没有找到与“{text}”直接相关的归档信息。"
    lines = [f"找到 {len(items)} 条相关信息："]
    for index, item in enumerate(items, start=1):
        lines.append(f"{index}. {item['title']}")
        if item.get("site_name") or item.get("section_name"):
            site_name = item.get("site_name") or "-"
            section_name = item.get("section_name") or "-"
            lines.append(f"   来源：{site_name} / {section_name}")
        if item.get("snippet"):
            lines.append(f"   命中：{item['snippet']}")
        if item.get("source_url"):
            lines.append(f"   原文：{item['source_url']}")
        if item.get("backend_url"):
            lines.append(f"   后台：{item['backend_url']}")
    return "\n".join(lines)


def latest_attachment_version(db: Session, attachment_id: int) -> AttachmentVersion | None:
    return db.scalar(
        select(AttachmentVersion)
        .where(AttachmentVersion.attachment_id == attachment_id)
        .order_by(AttachmentVersion.version_no.desc(), AttachmentVersion.id.desc())
    )


def get_or_create_document_text(
    db: Session,
    attachment: Attachment,
    version: AttachmentVersion | None,
) -> DocumentText:
    query = select(DocumentText).where(DocumentText.attachment_id == attachment.id)
    if version:
        query = query.where(DocumentText.attachment_version_id == version.id)
    else:
        query = query.where(DocumentText.attachment_version_id.is_(None))
    document_text = db.scalar(query)
    if document_text is None:
        document_text = DocumentText(
            attachment_id=attachment.id,
            attachment_version_id=version.id if version else None,
        )
        db.add(document_text)
        db.flush()
    return document_text


def get_or_create_search_index(db: Session, entity_type: str, entity_id: int) -> SearchIndex:
    item = db.scalar(
        select(SearchIndex)
        .where(SearchIndex.entity_type == entity_type)
        .where(SearchIndex.entity_id == entity_id)
    )
    if item is None:
        item = SearchIndex(entity_type=entity_type, entity_id=entity_id, title="")
        db.add(item)
        db.flush()
    return item


def resolve_attachment_file(
    settings: Settings,
    attachment: Attachment,
    version: AttachmentVersion | None,
) -> Path:
    local_path = version.local_path if version else attachment.local_path
    if not local_path:
        return settings.storage_root / "__missing__"
    path = Path(local_path)
    if path.is_absolute():
        return path
    return settings.storage_root / path


def parser_name_for_suffix(suffix: str) -> str:
    return {
        ".pdf": "pdf_text",
        ".docx": "docx_text",
        ".xlsx": "xlsx_text",
        ".xlsm": "xlsx_text",
    }.get(suffix.lower(), "unsupported")


def extract_text(path: Path, suffix: str) -> str:
    if suffix == ".pdf":
        return extract_pdf_text(path)
    if suffix == ".docx":
        return extract_docx_text(path)
    if suffix in {".xlsx", ".xlsm"}:
        return extract_xlsx_text(path)
    raise ValueError(f"unsupported file extension: {suffix}")


def extract_pdf_text(path: Path) -> str:
    reader = PdfReader(str(path))
    texts = [page.extract_text() or "" for page in reader.pages]
    return normalize_text("\n".join(texts))


def extract_docx_text(path: Path) -> str:
    document = Document(str(path))
    paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    table_cells: list[str] = []
    for table in document.tables:
        for row in table.rows:
            table_cells.extend(cell.text for cell in row.cells if cell.text.strip())
    return normalize_text("\n".join([*paragraphs, *table_cells]))


def extract_xlsx_text(path: Path) -> str:
    workbook = load_workbook(path, read_only=True, data_only=True)
    lines: list[str] = []
    try:
        for sheet in workbook.worksheets:
            lines.append(f"## {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                values = [
                    str(value).strip() for value in row if value is not None and str(value).strip()
                ]
                if values:
                    lines.append("\t".join(values))
    finally:
        workbook.close()
    return normalize_text("\n".join(lines))


def normalize_text(value: str) -> str:
    return "\n".join(
        line.strip() for line in value.replace("\r", "\n").splitlines() if line.strip()
    )


def sha256_text(value: str) -> str:
    return hashlib.sha256(value.encode("utf-8")).hexdigest()


def result_from_index(item: SearchIndex, terms: list[str]) -> KnowledgeSearchResult:
    body = item.body or ""
    match_fields: list[str] = []
    score = 0
    for term in terms:
        if term in item.title:
            score += 100
            if "title" not in match_fields:
                match_fields.append("title")
        if term in body:
            score += 20
            if "body" not in match_fields:
                match_fields.append("body")
        if item.site_name and term in item.site_name:
            score += 5
            if "site_name" not in match_fields:
                match_fields.append("site_name")
        if item.section_name and term in item.section_name:
            score += 5
            if "section_name" not in match_fields:
                match_fields.append("section_name")
    return KnowledgeSearchResult(
        entity_type=item.entity_type,
        entity_id=item.entity_id,
        title=item.title,
        snippet=make_snippet(body or item.title, terms),
        match_fields=match_fields,
        source_url=item.source_url,
        backend_path=item.backend_path,
        site_name=item.site_name,
        section_name=item.section_name,
        published_at=item.published_at,
        score=score,
    )


def make_snippet(text: str, terms: list[str], length: int = 160) -> str:
    if not text:
        return ""
    start = 0
    for term in terms:
        index = text.find(term)
        if index >= 0:
            start = max(0, index - length // 3)
            break
    snippet = text[start : start + length]
    prefix = "..." if start > 0 else ""
    suffix = "..." if start + length < len(text) else ""
    return f"{prefix}{snippet}{suffix}"


def search_result_to_dict(item: KnowledgeSearchResult, base_url: str | None = None) -> dict:
    backend_url = item.backend_path
    if base_url and item.backend_path:
        backend_url = f"{base_url.rstrip('/')}{item.backend_path}"
    return {
        "entity_type": item.entity_type,
        "entity_id": item.entity_id,
        "title": item.title,
        "snippet": item.snippet,
        "match_fields": item.match_fields,
        "source_url": item.source_url,
        "backend_url": backend_url,
        "site_name": item.site_name,
        "section_name": item.section_name,
        "published_at": item.published_at.isoformat() if item.published_at else None,
        "score": item.score,
    }


def kb_stats(db: Session) -> dict[str, int]:
    return {
        "document_texts": int(db.scalar(select(func.count(DocumentText.id))) or 0),
        "parsed_success": int(
            db.scalar(select(func.count(DocumentText.id)).where(DocumentText.status == "success"))
            or 0
        ),
        "parsed_failed": int(
            db.scalar(select(func.count(DocumentText.id)).where(DocumentText.status == "failed"))
            or 0
        ),
        "parsed_unsupported": int(
            db.scalar(
                select(func.count(DocumentText.id)).where(DocumentText.status == "unsupported")
            )
            or 0
        ),
        "search_index": int(db.scalar(select(func.count(SearchIndex.id))) or 0),
    }
